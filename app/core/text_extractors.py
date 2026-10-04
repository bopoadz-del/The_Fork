"""The registry of text extractors: the ONE declaration of ingestible formats.

A format is ingestible only when a working text extractor is registered for
it (owner ruling, docs/INGEST_EXCLUSION_RULE.md). Each extractor below
declares the formats it reads with ``@reads(...)``;
``ingest_status.TEXT_BEARING_EXTS`` is built from this registry, and every
door (project upload, ``/upload`` with a project, Drive ingest and import,
CDE ingest) asks it. To support a format, register an extractor here --
there is no other list to edit.

Compressed files (zip, rar and every other compressed-folder format) are
never registered: the platform never opens, unpacks or indexes one
(owner ruling; see ``app/core/compressed.py``).

Light by design: nothing heavy is imported at module level (``ingest_status``
imports this module). Each extractor imports what it needs when it runs.
An extractor takes ``(file_path, filename, force_ocr)`` and returns
``(text, meta)``; it may raise -- ``doc_index._extract_with_meta`` turns a
failure into ``extract_failed`` meta (and a memory failure into MemoryError).
"""
from __future__ import annotations

from dataclasses import dataclass
from typing import Any, Callable, Dict, FrozenSet, Optional, Tuple

ExtractFn = Callable[[str, str, bool], Tuple[str, Dict[str, Any]]]


@dataclass(frozen=True)
class Extractor:
    read: ExtractFn
    # An Office Open XML package: sniffed for a non-zip payload first
    # (doc_index._quarantine_meta) before the reader opens it.
    office_package: bool = False


_REGISTRY: Dict[str, Extractor] = {}


def reads(*exts: str, **options: Any) -> Callable[[ExtractFn], ExtractFn]:
    """Register the decorated function as the text extractor for ``exts``."""

    def register(fn: ExtractFn) -> ExtractFn:
        for ext in exts:
            if ext in _REGISTRY:
                raise ValueError(f"two extractors registered for {ext}")
            _REGISTRY[ext] = Extractor(read=fn, **options)
        return fn

    return register


def formats() -> FrozenSet[str]:
    """Every format a registered extractor reads (lower-case, with the dot)."""
    return frozenset(_REGISTRY)


def extractor_for(ext: str) -> Optional[Extractor]:
    return _REGISTRY.get((ext or "").lower())


# ── extractors ───────────────────────────────────────────────────────────────


@reads(".txt", ".md", ".csv", ".json", ".xml")
def _plain_text(file_path: str, filename: str, force_ocr: bool):
    from app.core import file_crypto

    return file_crypto.read_document(file_path).decode("utf-8", errors="replace"), {}


@reads(".pdf")
def _pdf(file_path: str, filename: str, force_ocr: bool):
    # Per-page: each page's text layer, OCR for image-only pages, table rows.
    from app.core import doc_index

    return doc_index._extract_pdf(file_path, filename, force_ocr=force_ocr)


@reads(".doc")
def _word_97(file_path: str, filename: str, force_ocr: bool):
    from app.core import doc_index

    return doc_index._extract_doc(file_path), {}


@reads(".docx", office_package=True)
def _docx(file_path: str, filename: str, force_ocr: bool):
    import docx

    from app.core import doc_index, file_crypto

    with file_crypto.open_plaintext(file_path) as readable_path:
        # python-docx `.paragraphs` excludes table cells; `.tables` is
        # top-level only and ignores nested w:tbl; text-boxes are invisible
        # to both. Letter signature blocks live in those structures.
        return doc_index._docx_plain_text(docx.Document(readable_path)), {}


def _sheet_rows(rows) -> str:
    """Row-wise, not cell-wise: a priced-BOQ line stays together
    ("<desc> | <qty> | <unit> | <rate> | <total>"). Flattening every cell into
    one blob separates a rate from its item and makes unit rates unfindable."""
    parts = []
    for row in rows:
        cells = [str(v).strip() for v in row if v is not None and str(v).strip()]
        if cells:
            parts.append(" | ".join(cells))
    return "\n".join(parts)


@reads(".xlsx", office_package=True)
def _xlsx(file_path: str, filename: str, force_ocr: bool):
    import openpyxl

    from app.core import file_crypto

    with file_crypto.open_plaintext(file_path) as readable_path:
        wb = openpyxl.load_workbook(readable_path, data_only=True)
        return "\n".join(
            _sheet_rows([cell.value for cell in row] for row in wb[name])
            for name in wb.sheetnames
        ), {}


@reads(".xls")
def _xls(file_path: str, filename: str, force_ocr: bool):
    # Legacy BIFF workbook: the same row-wise spreadsheet text as .xlsx,
    # read with xlrd (pure Python).
    import xlrd

    from app.core import file_crypto

    book = xlrd.open_workbook(file_contents=file_crypto.read_document(file_path))
    sheets = []
    for sheet in book.sheets():
        rows = (
            [_xls_cell(cell, book.datemode) for cell in sheet.row(r)]
            for r in range(sheet.nrows)
        )
        sheets.append(_sheet_rows(rows))
    return "\n".join(sheets), {}


def _xls_cell(cell, datemode: int):
    import xlrd

    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
        return None
    if cell.ctype == xlrd.XL_CELL_NUMBER and float(cell.value).is_integer():
        return int(cell.value)
    if cell.ctype == xlrd.XL_CELL_DATE:
        return xlrd.xldate.xldate_as_datetime(cell.value, datemode).date().isoformat()
    return cell.value


@reads(".pptx", office_package=True)
def _pptx(file_path: str, filename: str, force_ocr: bool):
    from app.core import doc_index

    return doc_index._extract_pptx_with_meta(file_path)


# PowerPoint 97-2003 record types ([MS-PPT] 2.13.24) that hold slide text.
_PPT_TEXT_CHARS_ATOM = 0x0FA0  # UTF-16LE
_PPT_TEXT_BYTES_ATOM = 0x0FA8  # 8-bit (low bytes of UTF-16)
_PPT_MAIN_MASTER = 0x03F8      # slide master: placeholder boilerplate only


@reads(".ppt")
def _powerpoint_97(file_path: str, filename: str, force_ocr: bool):
    # Legacy binary PowerPoint, read in pure Python: walk the record tree of
    # the "PowerPoint Document" stream and keep the text atoms (placeholders
    # and text boxes alike). catppt was tried first: it prints only master
    # theme names for files saved by current PowerPoint.
    import struct

    import olefile

    from app.core import file_crypto

    with file_crypto.open_plaintext(file_path) as readable_path:
        ole = olefile.OleFileIO(readable_path)
        try:
            stream = ole.openstream("PowerPoint Document").read()
        finally:
            ole.close()

    texts = []
    pos, end = 0, len(stream)
    while pos + 8 <= end:
        ver_inst, rec_type, rec_len = struct.unpack_from("<HHI", stream, pos)
        body = pos + 8
        if (ver_inst & 0x000F) == 0x000F:  # container
            # descend, except into a master's "Click to edit ..." prompts
            pos = body + rec_len if rec_type == _PPT_MAIN_MASTER else body
            continue
        if rec_type == _PPT_TEXT_CHARS_ATOM:
            texts.append(stream[body:body + rec_len].decode("utf-16-le", errors="replace"))
        elif rec_type == _PPT_TEXT_BYTES_ATOM:
            texts.append(stream[body:body + rec_len].decode("latin-1", errors="replace"))
        pos = body + rec_len
    lines = (t.replace("\r", "\n").strip() for t in texts)
    return "\n".join(t for t in lines if t), {}


@reads(".rtf")
def _rtf(file_path: str, filename: str, force_ocr: bool):
    from striprtf.striprtf import rtf_to_text

    from app.core import file_crypto

    raw = file_crypto.read_document(file_path).decode("cp1252", errors="replace")
    return rtf_to_text(raw, errors="ignore"), {}


@reads(".htm", ".html")
def _html(file_path: str, filename: str, force_ocr: bool):
    from bs4 import BeautifulSoup

    from app.core import file_crypto

    soup = BeautifulSoup(file_crypto.read_document(file_path), "html.parser")
    for node in soup(["script", "style", "noscript", "template"]):
        node.decompose()
    lines = (line.strip() for line in soup.get_text("\n").splitlines())
    return "\n".join(line for line in lines if line), {}


@reads(".msg")
def _outlook_msg(file_path: str, filename: str, force_ocr: bool):
    from app.core import doc_index

    return doc_index._extract_msg(file_path), {}


@reads(".ifc")
def _ifc(file_path: str, filename: str, force_ocr: bool):
    # The BIM element census (STEP census when the BIM extractor is absent).
    from app.core import doc_index

    return "\n".join(doc_index._ifc_census(file_path, filename)), {}
