"""Phase 2: chat / staged generate_wbs codes must equal the schedule xlsx.

Live FAIL on ~72fcfd2 / 7ace4df: a generate_wbs turn printed hierarchical
codes (1.1, 1.2, 2.1, …). Conversation message export matched that outline.
The Schedule Builder / ``export/schedule-from-brief`` workbook for the same
conversation wrote different WBS / activity codes (10, 109, 116, …) because
the download re-ran a template WBS instead of exporting the staged snapshot,
and the L2 WBS column stored phase slugs rather than the chat codes.

These tests pin the bind (no live LLM): the workbook ID and WBS columns must
be the conversation's staged codes, never a regenerated scaffold.
"""
from __future__ import annotations

import io

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.containers.construction import ConstructionContainer
from app.core.conversation_wbs import stage_conversation_wbs
from app.core.predefined_reasoning import format_wbs_outline
from app.main import app
from tests.conftest import requires_construction_kit

H = {"Authorization": "Bearer cb_dev_key"}
_XLSX_MEDIA = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

BRIEF = "Generate a high-level WBS for a 10-storey office tower."


@pytest.fixture
def client():
    with TestClient(app) as c:
        yield c


def _pid(client, name="WBS Code Match"):
    r = client.post("/v1/projects", json={"name": name, "client": "ACME"}, headers=H)
    assert r.status_code == 201, r.text
    return r.json()["id"]


def _l2_id_wbs(content: bytes) -> list[tuple[str, str]]:
    wb = openpyxl.load_workbook(io.BytesIO(content))
    assert "L2 Schedule" in wb.sheetnames
    ws = wb["L2 Schedule"]
    headers = [c.value for c in ws[3]]
    assert headers[0] == "ID" and headers[1] == "WBS"
    rows = []
    for row in ws.iter_rows(min_row=4, max_col=3, values_only=True):
        aid, wbs, name = row[0], row[1], row[2]
        if not name or name == "TOTAL" or str(name).startswith("Project Duration"):
            continue
        rows.append((str(aid), "" if wbs is None else str(wbs)))
    return rows


def _looks_like_wbs_code(value: str) -> bool:
    """Chat-style hierarchical code: 1 / 1.1 / 2.1.3 — not a phase slug."""
    raw = (value or "").strip()
    if not raw:
        return False
    parts = raw.split(".")
    return all(p.isdigit() for p in parts)


@requires_construction_kit
@pytest.mark.asyncio
async def test_generate_wbs_chat_codes_are_hierarchical_not_slugs():
    """The glass prints wbs_tree keys (1.1, 1.2, 2.1). Activities must carry
    the same codes so an export can write them instead of site_preparation."""
    wbs = await ConstructionContainer().generate_wbs(
        {"brief": BRIEF},
        {"target_count": 40, "project_type": "building", "start_date": "2026-03-01"},
    )
    tree = wbs["wbs_tree"]
    assert "1.1" in tree and "1.2" in tree and "2.1" in tree
    outline = format_wbs_outline(tree)
    assert "1.1 " in outline and "2.1 " in outline
    assert "site_preparation" not in outline

    acts = wbs["activities"]
    assert acts
    for a in acts:
        code = str(a.get("wbs_code") or a.get("wbs") or "")
        assert _looks_like_wbs_code(code), (
            f"activity {a.get('id')!r} has no hierarchical WBS code "
            f"(got wbs_code={a.get('wbs_code')!r} wbs={a.get('wbs')!r} "
            f"wbs_phase={a.get('wbs_phase')!r})"
        )
        assert code in tree, (
            f"activity {a.get('id')!r} code {code!r} is not a wbs_tree key"
        )


@requires_construction_kit
@pytest.mark.asyncio
async def test_schedule_from_brief_xlsx_codes_match_staged_chat_wbs(
    client, tmp_path, monkeypatch,
):
    """Same conversation_id: schedule-from-brief must write the staged codes.

    A conflicting brief / target_count that would rebuild a 200-activity
    data-center template must not replace 1.1 / 1.2 / 2.1 with 10 / 109 / 116
    (or any other regenerated numbering).
    """
    monkeypatch.setenv("DATA_DIR", str(tmp_path))
    pid = _pid(client)
    cid = f"ws-{pid}-1770000001000"
    wbs = await ConstructionContainer().generate_wbs(
        {"brief": BRIEF},
        {"target_count": 40, "project_type": "building", "start_date": "2026-03-01"},
    )
    snap = stage_conversation_wbs(cid, wbs)
    assert snap is not None
    staged_ids = [str(a["id"]) for a in snap["activities"]]
    staged_tree = {str(k) for k in (snap.get("wbs_tree") or {})}
    assert "1.1" in staged_tree and "2.1" in staged_tree

    r = client.post(
        f"/v1/projects/{pid}/export/schedule-from-brief",
        json={
            "brief": "hyperscale data center campus",
            "target_count": 200,
            "project_type": "data_center",
            "conversation_id": cid,
        },
        headers=H,
    )
    assert r.status_code == 200, r.text
    assert r.headers.get("X-WBS-Source") == "conversation"
    rows = _l2_id_wbs(r.content)
    assert [aid for aid, _ in rows] == staged_ids
    xlsx_wbs = [w for _, w in rows]
    assert xlsx_wbs
    assert all(_looks_like_wbs_code(w) for w in xlsx_wbs), (
        f"xlsx WBS column must be the chat/staged codes, not phase slugs "
        f"or sequential ids; got {sorted(set(xlsx_wbs))[:12]}"
    )
    assert all(w in staged_tree for w in xlsx_wbs), (
        f"xlsx WBS codes {sorted(set(xlsx_wbs))[:12]} are not in the "
        f"staged wbs_tree {sorted(staged_tree)[:12]}"
    )
    assert {"1.1", "1.2", "2.1"} <= set(xlsx_wbs)
    assert not any(w.isdigit() and int(w) >= 10 for w in xlsx_wbs), (
        f"regenerated sequential codes leaked into the workbook: {xlsx_wbs[:8]}"
    )
    assert "site_preparation" not in xlsx_wbs
    assert "Hall" not in " ".join(str(n) for n in xlsx_wbs)


@requires_construction_kit
def test_schedule_from_brief_with_conversation_id_does_not_regenerate(
    client, tmp_path, monkeypatch,
):
    """conversation_id set but nothing staged → honest 404, not a template."""
    monkeypatch.setenv("DATA_DIR", str(tmp_path))
    pid = _pid(client, "No Staged WBS")
    cid = f"ws-{pid}-1770000002000"
    r = client.post(
        f"/v1/projects/{pid}/export/schedule-from-brief",
        json={
            "brief": "hyperscale data center campus",
            "target_count": 200,
            "project_type": "data_center",
            "conversation_id": cid,
        },
        headers=H,
    )
    assert r.status_code == 404, r.text
    assert "template" in r.json()["detail"].lower()
