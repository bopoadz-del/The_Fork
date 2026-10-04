"""The RAG takes text formats only, at every door (docs/INGEST_EXCLUSION_RULE.md).

docs/RAG_GAPS_REVIEW_2026-09-12.md section E: CAD, images, Google Earth,
video, GIS internals and fonts are excluded BY DESIGN. Live 2026-10-04 a
Drive ingest downloaded and registered 8 videos, hundreds of drawings and
photos anyway. These tests pin the rule where files enter:
  * one declaration (ingest_status.TEXT_BEARING_EXTS) the indexer derives from;
  * archive members that are not text are never read;
  * a project upload / Drive import of an excluded format is refused (415)
    and writes nothing;
  * a chat photo is still accepted as question context and creates no
    document and no chunks.
The Drive ingest discovery filter is covered in test_p1b_ingest_accounting.
"""
from __future__ import annotations

import io
import zipfile
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from app.core import ingest_status as ist
from app.main import app

H = {"Authorization": "Bearer cb_dev_key"}


@pytest.fixture(scope="module")
def client():
    with TestClient(app) as c:
        yield c


def _new_project(client, name="Text Only"):
    r = client.post("/v1/projects", json={"name": name, "client": "ACME"}, headers=H)
    assert r.status_code == 201, r.text
    return r.json()


def _doc_count(client, pid):
    r = client.get(f"/v1/projects/{pid}/documents", headers=H)
    assert r.status_code == 200, r.text
    body = r.json()
    return len(body.get("documents", body) if isinstance(body, dict) else body)


# ── one declaration ────────────────────────────────────────────────────────


def test_the_indexer_derives_its_formats_from_the_single_declaration():
    from app.core import doc_index
    from app.routers import projects

    # Equality, not identity: other tests reload these modules.
    assert set(doc_index._SUPPORTED_EXTS) == set(ist.TEXT_BEARING_EXTS)
    assert set(projects.ALLOWED_DOC_EXTENSIONS) == set(ist.TEXT_BEARING_EXTS)


@pytest.mark.parametrize("name", ["drone.mp4", "plan.dwg", "site.jpg", "earth.kmz", "font.ttf", "pack.zip"])
def test_excluded_formats_are_not_ingestible(name):
    assert not ist.is_ingestible(name)


@pytest.mark.parametrize("name", ["contract.pdf", "boq.xlsx", "notes.txt", "Spec.DOCX"])
def test_text_formats_are_ingestible(name):
    assert ist.is_ingestible(name)


# ── upload doors ───────────────────────────────────────────────────────────


@pytest.mark.parametrize("name,mime", [
    ("drone.mp4", "video/mp4"),
    ("plan.dwg", "application/acad"),
    ("site.jpg", "image/jpeg"),
])
def test_project_upload_of_an_excluded_format_is_refused_and_writes_nothing(client, name, mime):
    proj = _new_project(client, f"Refuse {name}")
    before = _doc_count(client, proj["id"])
    r = client.post(
        f"/v1/projects/{proj['id']}/documents",
        files={"file": (name, b"\x00" * 128, mime)}, headers=H,
    )
    assert r.status_code == 415, r.text
    assert "text formats only" in r.json()["detail"]
    assert _doc_count(client, proj["id"]) == before


def test_upload_with_a_project_refuses_an_excluded_format(client):
    proj = _new_project(client, "Refuse via /upload")
    before = _doc_count(client, proj["id"])
    r = client.post(
        "/v1/upload",
        files={"file": ("drone.mp4", b"\x00" * 128, "video/mp4")},
        data={"project_id": proj["id"]}, headers=H,
    )
    assert r.status_code == 415, r.text
    assert "text formats only" in r.json()["detail"]
    assert _doc_count(client, proj["id"]) == before


def test_a_text_upload_to_a_project_still_lands(client):
    proj = _new_project(client, "Text lands")
    r = client.post(
        f"/v1/projects/{proj['id']}/documents",
        files={"file": ("notes.txt", b"Concrete cover 50 mm to footings.", "text/plain")},
        headers=H,
    )
    assert r.status_code in (200, 201), r.text


# ── chat photos stay question context ──────────────────────────────────────


def test_chat_photo_is_accepted_as_context_and_creates_no_document(client, monkeypatch):
    from app import dependencies

    class _Image:
        async def execute(self, inputs, params):
            return {"result": {"safety_qaqc": [], "summary_by_class": {}}}

    monkeypatch.setattr(dependencies, "get_block_instance", lambda name: _Image())

    from app.core.db import SessionLocal
    from app.core.models import Document

    with SessionLocal() as s:
        docs_before = s.query(Document).count()
    r = client.post(
        "/v1/chat/analyze-photo",
        files={"file": ("site.jpg", b"\xff\xd8\xff" + b"\x00" * 256, "image/jpeg")},
        headers=H,
    )
    assert r.status_code == 200, r.text
    with SessionLocal() as s:
        assert s.query(Document).count() == docs_before


# ── deleting a document removes everything it left behind ──────────────────


def test_purge_document_removes_the_stored_file_the_row_and_the_index_entry(tmp_path, monkeypatch):
    import importlib

    monkeypatch.setenv("DATA_DIR", str(tmp_path))
    monkeypatch.delenv("DATA_ENCRYPTION_KEY", raising=False)
    from app.core import projects as projects_mod

    monkeypatch.setattr(projects_mod, "_initialized", False)
    projects_mod.init_db()
    from app.core import doc_index, file_crypto

    importlib.reload(doc_index)
    pid = projects_mod.create_project("Purge")["id"]
    path = str(tmp_path / "notes.txt")
    file_crypto.write_document(path, b"Concrete cover 50 mm to footings. " * 20)
    doc = projects_mod.add_document(pid, "notes.txt", file_path=path, size=600)
    doc_index._upsert_index_entry(pid, doc["id"], "document",
                                  {"document_id": doc["id"], "filename": "notes.txt", "chunk_count": 1})

    out = projects_mod.purge_document(doc["id"])

    assert out == {"deleted": True, "file_removed": True, "index_pruned": True}
    assert not (tmp_path / "notes.txt").exists()
    assert projects_mod.get_document(doc["id"]) is None
    assert (doc_index._load_index(pid) or {}).get("documents", []) == []


# ── F-FINISH step 2: every registered format reads its sample ─────────────

FIXTURES = Path(__file__).parent / "fixtures"
PROBE = "reinforced concrete pour sequence"


def _sample(ext: str, tmp_path) -> tuple[Path, str]:
    """A small real file of ``ext`` and the text its chunks must contain."""
    import json as _json

    path = tmp_path / f"sample{ext}"
    line = f"FORMATPROBE {ext}: {PROBE} for level three."
    if ext in {".txt", ".md", ".csv"}:
        path.write_text(line, encoding="utf-8")
    elif ext == ".json":
        path.write_text(_json.dumps({"note": line}), encoding="utf-8")
    elif ext == ".xml":
        path.write_text(f"<note>{line}</note>", encoding="utf-8")
    elif ext in {".htm", ".html"}:
        path.write_text(f"<html><script>var x=1;</script><body><p>{line}</p></body></html>",
                        encoding="utf-8")
    elif ext == ".pdf":
        import fitz

        pdf = fitz.open()
        pdf.new_page().insert_text((72, 72), line)
        pdf.save(str(path))
    elif ext == ".docx":
        import docx

        d = docx.Document()
        d.add_paragraph(line)
        d.save(str(path))
    elif ext == ".xlsx":
        import openpyxl

        wb = openpyxl.Workbook()
        wb.active.append(["FORMATPROBE", PROBE, 42])
        wb.save(str(path))
    elif ext == ".pptx":
        from pptx import Presentation
        from pptx.util import Inches

        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text_frame.text = line
        deck.save(str(path))
    elif ext == ".ifc":
        return FIXTURES / "sample_office.ifc", "IFC MODEL"
    else:  # legacy binaries written by the real applications
        return FIXTURES / "formats" / f"sample{ext}", PROBE
    return path, PROBE


@pytest.mark.parametrize("ext", sorted(ist.TEXT_BEARING_EXTS))
def test_every_registered_format_reads_its_sample_into_a_chunk(ext, tmp_path):
    """A format is ingestible only if a registered extractor reads it: each
    one, generated from the registry, yields a chunk holding its known text."""
    from app.core import doc_index

    path, expected = _sample(ext, tmp_path)
    assert path.exists(), f"no sample for registered format {ext}"
    text, meta = doc_index._extract_with_meta(str(path), path.name)
    chunks = doc_index.chunk_extracted_document(
        text, chunker=doc_index._chunker_for_document(path.name), filename=path.name,
    )
    assert any(expected.lower() in c.lower() for c in chunks), (ext, meta, text[:200])


# ── F-FINISH step 3: compressed files are never opened ────────────────────


def _zip_bytes() -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("spec.txt", "Retaining wall drainage clause seven.")
    return buf.getvalue()


def test_a_compressed_file_is_never_ingested_opened_or_accepted(client, tmp_path):
    from app.core import compressed, text_extractors
    from scripts.p1b_ingest_drive_server import _ingest_file

    blob = _zip_bytes()
    # No registered extractor reads a compressed folder.
    assert not {".zip", ".rar", ".7z", ".gz", ".tar"} & text_extractors.formats()
    assert not ist.is_ingestible("pack.zip")
    # Ingest: excluded before download.
    downloads: list = []

    class _Drive:
        def download_file_bytes(self, fid):
            downloads.append(fid)
            return blob, None

    _rel, result = _ingest_file(
        {"id": "drive-z", "name": "pack.zip", "_drive_path": "T/pack.zip",
         "mimeType": "application/zip", "size": len(blob)},
        "proj", tmp_path, "run-z", _Drive(),
    )
    assert downloads == [] and result["status"] == "error"
    # Upload, both routes, with and without a project.
    proj = _new_project(client, "Compressed")
    for name in ("pack.zip", "pack"):
        files = {"file": (name, blob, "application/zip")}
        for r in (
            client.post(f"/v1/projects/{proj['id']}/documents", files=files, headers=H),
            client.post("/upload", files=files, data={"project_id": proj["id"]}, headers=H),
            client.post("/upload", files=files, headers=H),
        ):
            assert r.status_code == 415, r.text
            assert r.json()["detail"] == compressed.COMPRESSED_UPLOAD_DETAIL
    assert _doc_count(client, proj["id"]) == 0


# ── F-FINISH step 4: a file with no text is never recorded as indexed ─────


def test_a_file_with_no_text_is_never_recorded_as_indexed(tmp_path, monkeypatch):
    import importlib

    import sqlalchemy as sa

    monkeypatch.setenv("DATA_DIR", str(tmp_path))
    monkeypatch.setenv("RAG_EMBEDDING_MODEL", "fake")
    monkeypatch.delenv("DATABASE_URL", raising=False)
    import app.core.db as db_mod
    from app.core import projects as projects_mod

    importlib.reload(db_mod)
    projects_mod = importlib.reload(projects_mod)
    projects_mod._initialized = False
    from app.core import doc_index

    doc_index = importlib.reload(doc_index)
    projects_mod.init_db()

    pid = projects_mod.create_project("Empty text")["id"]
    blank = tmp_path / "blank.txt"
    blank.write_text("   \n\n   ", encoding="utf-8")
    doc = projects_mod.add_document(pid, "blank.txt", file_path=str(blank), size=9)

    result = doc_index.index_project(pid)

    assert result["indexed"] == 0
    assert doc_index._load_index(pid)["documents"] == []
    assert projects_mod.get_document(doc["id"])["ingest_status"] == ist.ZERO_CHUNK

    # Rows already recorded INDEXED with no text are corrected, idempotently.
    stale = projects_mod.add_document(pid, "old.txt", file_path=str(blank), size=9)
    with db_mod.get_engine().begin() as conn:
        conn.execute(sa.text("UPDATE documents SET ingest_status = :s, chunk_count = 0 "
                             "WHERE id = :i"), {"s": ist.INDEXED, "i": stale["id"]})
        assert ist.correct_empty_indexed(conn) == 1
        assert ist.correct_empty_indexed(conn) == 0
    assert projects_mod.get_document(stale["id"])["ingest_status"] == ist.ZERO_CHUNK
