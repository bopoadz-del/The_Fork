"""Do not advance extractor_version when text extracted but embed failed.

``index_document`` treats RAG as best-effort: it logs RAG_INDEX_FAILED, sets
``rag_error``, and still returns status=ok. ``_stamp_index_ledger`` then
stamped EXTRACTOR_VERSION regardless. Re-extract trusted that stamp, so ~27
TEXT_SPARSE .docx were silently closed.

Text-extracted-but-not-embedded is open, not done. The extractor version
must stay where it was so ``docx_stale_extractor_open`` still selects the
row. Happy-path embed still advances the version / INDEXED as today.
"""
from __future__ import annotations

import hashlib
import importlib
import io
from pathlib import Path

import docx

from app.core.ingest_status import EXTRACTOR_VERSION, INDEXED, TEXT_SPARSE
from app.core import ingest_status as ist


DRIVE_STALE_ID = "driveStamp01"
EMBED_FAILED_SENTINEL = f"{EXTRACTOR_VERSION}/embed-failed"


def _reload(monkeypatch, tmp_path):
    monkeypatch.setenv("DATA_DIR", str(tmp_path))
    monkeypatch.setenv("RAG_EMBEDDING_MODEL", "fake")
    monkeypatch.delenv("DATABASE_URL", raising=False)
    import app.core.db as db_mod
    import app.core.users as users_mod
    from app.core import projects
    from app.core.rag.embeddings import reset_embedder_cache
    from app.core.rag.vector_store import reset_store_cache

    importlib.reload(db_mod)
    importlib.reload(users_mod)
    users_mod._initialized = False
    projects._initialized = False
    reset_store_cache()
    reset_embedder_cache()
    import app.core.doc_index as doc_index_mod
    importlib.reload(doc_index_mod)
    return importlib.reload(projects), users_mod


def _rich_docx_bytes() -> bytes:
    document = docx.Document()
    document.add_paragraph(" ".join(f"word{i}" for i in range(800)))
    document.add_paragraph(" ".join(f"extra{i}" for i in range(200)))
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def _thin_docx_bytes() -> bytes:
    document = docx.Document()
    document.add_paragraph("one window only")
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def _seed_project(projects, users):
    projects.init_db()
    users.ensure_user_exists("u1")
    projects.create_project(name="P", client="C", user_id="u1")
    return projects.list_projects("u1")[0]


def _add_docx(
    projects,
    project_id: str,
    tmp_path: Path,
    *,
    name: str,
    ingest_status: str,
    extractor_version: str | None,
    chunk_count: int,
    raw: bytes | None = None,
    drive_file_id: str = DRIVE_STALE_ID,
    r2_object_key: str | None = "projects/p/stale.docx",
):
    dest = tmp_path / name
    payload = raw if raw is not None else _thin_docx_bytes()
    dest.write_bytes(payload)
    doc = projects.add_document(
        project_id=project_id,
        original_name=name,
        stored_as=name,
        file_path=str(dest),
        size=len(payload),
        content_sha256=hashlib.sha256(name.encode()).hexdigest(),
        metadata={
            "drive_file_id": drive_file_id,
            "r2_object_key": r2_object_key,
            "r2_bucket": "corpus",
        },
    )
    projects.stamp_document_index(
        doc["id"],
        chunk_count=chunk_count,
        ingest_status=ingest_status,
        ingest_status_reason="single_window:terminal",
        extractor_version=extractor_version,
    )
    return projects.get_document(doc["id"])


def _ledger_entry(project_id: str, document_id: str) -> dict:
    from app.core import doc_index

    saved = doc_index._load_index(project_id) or {}
    for entry in saved.get("documents") or []:
        if entry.get("document_id") == document_id:
            return entry
    return {}


def test_reextract_embed_fail_stays_stale_and_does_not_advance_version(
    monkeypatch, tmp_path,
):
    """Re-extract with embed stubbed to fail: stay stale-selectable."""
    projects, users = _reload(monkeypatch, tmp_path)
    proj = _seed_project(projects, users)
    stale = _add_docx(
        projects, proj["id"], tmp_path,
        name="stale.docx",
        ingest_status=TEXT_SPARSE,
        extractor_version="pre-sdt",
        chunk_count=1,
    )
    thin = _thin_docx_bytes()
    monkeypatch.setattr(
        "app.core.r2_storage.fetch_object_bytes",
        lambda key, bucket=None: thin if key == "projects/p/stale.docx" else None,
    )
    from app.core.rag import retriever as _rag

    monkeypatch.setattr(
        _rag, "index_chunks",
        lambda *_a, **_kw: (_ for _ in ()).throw(RuntimeError("403 Forbidden")),
    )

    from scripts.reextract_stale_docx import main

    assert main([]) == 0
    after = projects.get_document(stale["id"])
    assert after["extractor_version"] == "pre-sdt"
    assert after["extractor_version"] != EXTRACTOR_VERSION
    assert ist.docx_stale_extractor_open(
        after["ingest_status"],
        extension=ist.document_extension(after),
        extractor_version=after["extractor_version"],
    )
    result_entry = _ledger_entry(proj["id"], stale["id"])
    assert result_entry.get("rag_error")
    assert "403" in result_entry["rag_error"]


def test_index_document_embed_fail_does_not_stamp_current_version(
    monkeypatch, tmp_path,
):
    """First ingest: version stays unset so the row remains open."""
    projects, users = _reload(monkeypatch, tmp_path)
    proj = _seed_project(projects, users)
    dest = tmp_path / "fresh.docx"
    dest.write_bytes(_thin_docx_bytes())
    doc = projects.add_document(
        project_id=proj["id"],
        original_name="fresh.docx",
        stored_as="fresh.docx",
        file_path=str(dest),
        size=dest.stat().st_size,
        content_sha256="ab" * 32,
    )
    assert doc.get("extractor_version") is None

    from app.core import doc_index
    from app.core.rag import retriever as _rag

    monkeypatch.setattr(_rag, "index_chunks", lambda *_a, **_kw: 0)
    result = doc_index.index_document(proj["id"], doc["id"])
    assert result.get("rag_error")
    assert result.get("rag_indexed", 0) == 0

    after = projects.get_document(doc["id"])
    assert after["extractor_version"] != EXTRACTOR_VERSION
    assert after["extractor_version"] is None
    assert ist.docx_stale_extractor_open(
        after["ingest_status"],
        extension=".docx",
        extractor_version=after["extractor_version"],
    )
    entry = _ledger_entry(proj["id"], doc["id"])
    assert entry.get("rag_error")
    assert "0 of" in entry["rag_error"]


def test_embed_success_advances_version_and_indexes(monkeypatch, tmp_path):
    """Happy path unchanged: version advances, INDEXED when text is rich."""
    projects, users = _reload(monkeypatch, tmp_path)
    proj = _seed_project(projects, users)
    rich = _rich_docx_bytes()
    stale = _add_docx(
        projects, proj["id"], tmp_path,
        name="stale.docx",
        ingest_status=TEXT_SPARSE,
        extractor_version="pre-sdt",
        chunk_count=1,
        raw=rich,
    )
    monkeypatch.setattr(
        "app.core.r2_storage.fetch_object_bytes",
        lambda key, bucket=None: rich if key == "projects/p/stale.docx" else None,
    )

    from scripts.reextract_stale_docx import main

    assert main([]) == 0
    after = projects.get_document(stale["id"])
    assert after["ingest_status"] == INDEXED
    assert after["extractor_version"] == EXTRACTOR_VERSION
    assert after["chunk_count"] > 1
    assert not ist.docx_stale_extractor_open(
        after["ingest_status"],
        extension=".docx",
        extractor_version=after["extractor_version"],
    )
    entry = _ledger_entry(proj["id"], stale["id"])
    assert "rag_error" not in entry
    assert int(entry.get("rag_indexed") or 0) > 0


def test_reextract_sentinel_thin_embed_ok_terminal_closes(
    monkeypatch, tmp_path, capsys,
):
    """Sentinel + genuine thin + embed OK → stamp EXTRACTOR_VERSION.

    #575 left ``{EXTRACTOR_VERSION}/embed-failed`` in place after a
    1-chunk TEXT_SPARSE embed so the row could not false-close. That
    is now an infinite ``docx_stale_extractor_open`` loop: the extract
    succeeded, the vector landed, the sheet is honestly thin.
    Terminal-close the sentinel only in this case.
    """
    projects, users = _reload(monkeypatch, tmp_path)
    proj = _seed_project(projects, users)
    stale = _add_docx(
        projects, proj["id"], tmp_path,
        name="stale.docx",
        ingest_status=TEXT_SPARSE,
        extractor_version=EMBED_FAILED_SENTINEL,
        chunk_count=1,
    )
    thin = _thin_docx_bytes()
    monkeypatch.setattr(
        "app.core.r2_storage.fetch_object_bytes",
        lambda key, bucket=None: thin if key == "projects/p/stale.docx" else None,
    )

    from scripts.reextract_stale_docx import main

    assert main([]) == 0
    out = capsys.readouterr().out
    after = projects.get_document(stale["id"])
    assert after["ingest_status"] == TEXT_SPARSE
    assert after["chunk_count"] == 1
    assert after["extractor_version"] == EXTRACTOR_VERSION
    assert after["extractor_version"] != EMBED_FAILED_SENTINEL
    assert not ist.docx_stale_extractor_open(
        after["ingest_status"],
        extension=ist.document_extension(after),
        extractor_version=after["extractor_version"],
    )
    entry = _ledger_entry(proj["id"], stale["id"])
    assert "rag_error" not in entry
    assert int(entry.get("rag_indexed") or 0) == 1
    assert f"before={TEXT_SPARSE}/{EMBED_FAILED_SENTINEL}/1" in out
    assert f"after={TEXT_SPARSE}/{EXTRACTOR_VERSION}/1" in out


def test_reextract_sentinel_rag_error_does_not_advance(monkeypatch, tmp_path):
    """Sentinel + rag_error must stay open. Never terminal-close a miss."""
    projects, users = _reload(monkeypatch, tmp_path)
    proj = _seed_project(projects, users)
    stale = _add_docx(
        projects, proj["id"], tmp_path,
        name="stale.docx",
        ingest_status=TEXT_SPARSE,
        extractor_version=EMBED_FAILED_SENTINEL,
        chunk_count=1,
    )
    thin = _thin_docx_bytes()
    monkeypatch.setattr(
        "app.core.r2_storage.fetch_object_bytes",
        lambda key, bucket=None: thin if key == "projects/p/stale.docx" else None,
    )
    from app.core.rag import retriever as _rag

    monkeypatch.setattr(
        _rag, "index_chunks",
        lambda *_a, **_kw: (_ for _ in ()).throw(RuntimeError("403 Forbidden")),
    )

    from scripts.reextract_stale_docx import main

    assert main([]) == 0
    after = projects.get_document(stale["id"])
    assert after["extractor_version"] == EMBED_FAILED_SENTINEL
    assert after["extractor_version"] != EXTRACTOR_VERSION
    assert ist.docx_stale_extractor_open(
        after["ingest_status"],
        extension=ist.document_extension(after),
        extractor_version=after["extractor_version"],
    )
    entry = _ledger_entry(proj["id"], stale["id"])
    assert entry.get("rag_error")
    assert "403" in entry["rag_error"]


def test_index_document_sentinel_thin_embed_ok_terminal_closes(
    monkeypatch, tmp_path,
):
    """index_document: sentinel + TEXT_SPARSE + rag_indexed>0 advances."""
    projects, users = _reload(monkeypatch, tmp_path)
    proj = _seed_project(projects, users)
    dest = tmp_path / "thin.docx"
    dest.write_bytes(_thin_docx_bytes())
    doc = projects.add_document(
        project_id=proj["id"],
        original_name="thin.docx",
        stored_as="thin.docx",
        file_path=str(dest),
        size=dest.stat().st_size,
        content_sha256="cd" * 32,
    )
    projects.stamp_document_index(
        doc["id"],
        chunk_count=1,
        ingest_status=TEXT_SPARSE,
        ingest_status_reason="single_window:terminal",
        extractor_version=EMBED_FAILED_SENTINEL,
    )

    from app.core import doc_index

    result = doc_index.index_document(proj["id"], doc["id"], stamp_as_indexed=False)
    assert result.get("status") == "ok"
    assert "rag_error" not in result
    assert int(result.get("rag_indexed") or 0) > 0

    after = projects.get_document(doc["id"])
    assert after["ingest_status"] == TEXT_SPARSE
    assert after["extractor_version"] == EXTRACTOR_VERSION
    assert after["extractor_version"] != EMBED_FAILED_SENTINEL
    assert not ist.docx_stale_extractor_open(
        after["ingest_status"],
        extension=".docx",
        extractor_version=after["extractor_version"],
    )


def test_index_document_nonsentinel_text_sparse_does_not_advance(
    monkeypatch, tmp_path,
):
    """#575 protection: plain TEXT_SPARSE + rag_indexed=1 stays unstamped.

    The pre-#575 false-close class (never embed-failed) must not close
    here. Part B reopens only the large under-extract subset.
    """
    projects, users = _reload(monkeypatch, tmp_path)
    proj = _seed_project(projects, users)
    dest = tmp_path / "thin.docx"
    dest.write_bytes(_thin_docx_bytes())
    doc = projects.add_document(
        project_id=proj["id"],
        original_name="thin.docx",
        stored_as="thin.docx",
        file_path=str(dest),
        size=dest.stat().st_size,
        content_sha256="ef" * 32,
    )
    projects.stamp_document_index(
        doc["id"],
        chunk_count=1,
        ingest_status=TEXT_SPARSE,
        ingest_status_reason="single_window:terminal",
        extractor_version="pre-sdt",
    )

    from app.core import doc_index

    result = doc_index.index_document(proj["id"], doc["id"], stamp_as_indexed=False)
    assert result.get("status") == "ok"
    assert "rag_error" not in result
    assert int(result.get("rag_indexed") or 0) == 1

    after = projects.get_document(doc["id"])
    assert after["ingest_status"] == TEXT_SPARSE
    assert after["extractor_version"] == "pre-sdt"
    assert after["extractor_version"] != EXTRACTOR_VERSION
    assert ist.docx_stale_extractor_open(
        after["ingest_status"],
        extension=".docx",
        extractor_version=after["extractor_version"],
    )
