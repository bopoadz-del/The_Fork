"""R2: census compares stored chars to the fixed extractor on fixtures."""
from __future__ import annotations

import importlib
from pathlib import Path

import docx

from tests.test_docx_content_control_extraction import _sdt_block, _sdt_inline


def _reload(monkeypatch, tmp_path):
    monkeypatch.setenv("DATA_DIR", str(tmp_path))
    monkeypatch.setenv("RAG_EMBEDDING_MODEL", "fake")
    monkeypatch.delenv("DATABASE_URL", raising=False)
    import app.core.db as db_mod
    import app.core.users as users_mod
    from app.core import projects
    from app.core.rag.vector_store import reset_store_cache
    from app.core.rag.embeddings import reset_embedder_cache

    importlib.reload(db_mod)
    importlib.reload(users_mod)
    users_mod._initialized = False
    projects._initialized = False
    reset_store_cache()
    reset_embedder_cache()
    import app.core.doc_index as doc_index_mod
    importlib.reload(doc_index_mod)
    return importlib.reload(projects), users_mod


def _letter_docx(path: Path) -> Path:
    document = docx.Document()
    document.add_paragraph("Dear Sir or Madam,")
    document.add_paragraph("The works described above are complete.")
    document.add_paragraph("Yours sincerely,")
    document.element.body.append(
        _sdt_block(["Regards", "A Signatory", "Engineer's Representative"])
    )
    subject = document.add_paragraph("Ref. No.: ")
    subject._p.append(_sdt_inline("REF-000372"))
    subject.add_run(" — completion notice")
    document.save(path)
    return path


def test_census_delta_zero_when_source_matches_fixed_extract(monkeypatch, tmp_path):
    projects, users = _reload(monkeypatch, tmp_path)
    projects.init_db()
    users.ensure_user_exists("u1")
    projects.create_project(name="P", client="C", user_id="u1")
    proj = projects.list_projects("u1")[0]
    src = _letter_docx(tmp_path / "letter.docx")
    sha = __import__("hashlib").sha256(src.read_bytes()).hexdigest()
    doc = projects.add_document(
        project_id=proj["id"],
        original_name="letter.docx",
        stored_as="letter.docx",
        file_path=str(src),
        size=src.stat().st_size,
        content_sha256=sha,
    )
    from app.core import doc_index

    doc_index.index_document(proj["id"], doc["id"], stamp_as_indexed=True)

    from scripts.extraction_census import run_census

    out = tmp_path / "EXTRACTION_CENSUS.md"
    report = run_census(project_id=proj["id"], out_path=out)
    rows = report["rows"]
    assert len(rows) == 1
    assert rows[0]["source_available"] == "Y"
    assert rows[0]["delta"] == 0
    assert "chars stored" in report["markdown"]
    assert out.is_file()


def test_census_then_reingest_brings_delta_to_zero(monkeypatch, tmp_path):
    projects, users = _reload(monkeypatch, tmp_path)
    projects.init_db()
    users.ensure_user_exists("u1")
    projects.create_project(name="P", client="C", user_id="u1")
    proj = projects.list_projects("u1")[0]
    src = _letter_docx(tmp_path / "letter.docx")
    sha = __import__("hashlib").sha256(src.read_bytes()).hexdigest()
    doc = projects.add_document(
        project_id=proj["id"],
        original_name="letter.docx",
        stored_as="letter.docx",
        file_path=str(src),
        size=src.stat().st_size,
        content_sha256=sha,
    )
    from app.core.rag.embeddings import get_embedder
    from app.core.rag.vector_store import get_store

    short = "Dear Sir or Madam, The works are complete. Yours sincerely,"
    store = get_store()
    store.upsert_chunks(
        proj["id"], doc["id"], [short], get_embedder().encode([short]),
    )

    from scripts.extraction_census import apply_reingest, census_row, run_census

    row = census_row(projects.get_document(doc["id"]))
    assert row["source_available"] == "Y"
    assert row["delta"] > 0

    applied = apply_reingest(row)
    assert applied.get("skipped") is False
    new_id = applied["new_id"]
    hidden = projects.get_document(doc["id"])
    assert hidden["retrieval_visible"] is False
    assert hidden["superseded_by"] == new_id

    new_row = census_row(projects.get_document(new_id))
    assert new_row["source_available"] == "Y"
    assert new_row["delta"] == 0
