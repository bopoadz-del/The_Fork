"""F-BAT-D H2: export the chat-built F1 WBS, not a 204-activity scaffold.

Live leftover on BASELINE 0d9fd23 (Master Corpus):

    After F1 ("Schedule built: 2 activities…"), "Export F1 WBS as xlsx"
    returned the building-template scaffold (204 activities). Per-answer
    Excel download 404'd "No assistant messages in this conversation."
    The Schedule (Excel) button downloaded Master_Corpus_schedule.xlsx
    — a valid workbook with a Dur column, but the generic scaffold.

These tests pin the bind (no live LLM). A BOQ-derived WBS that was
built in the conversation must be the xlsx; a missing snapshot is an
honest failure, never a silent 204-activity success.
"""
from __future__ import annotations

import asyncio
import io
import json

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.core.conversation_wbs import (
    H2_EXPORT_ASK,
    clear_conversation_wbs,
    conversation_schedule_export_descriptor,
    fulfill_wbs_export,
    load_conversation_wbs,
    message_looks_like_wbs_answer,
    message_wants_wbs_export,
    refuse_scaffold_for_boq_wbs_ask,
    snapshot_from_wbs,
    stage_conversation_wbs,
)
from app.core.predefined_reasoning import message_wants_boq_scope_wbs
from app.lib.boq_schedule import (
    activities_from_boq_scope_outline,
    wbs_tree_from_boq_items,
)
from app.main import app
from tests.conftest import requires_construction_kit

H = {"Authorization": "Bearer cb_dev_key"}
_XLSX_MEDIA = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

F1_ASK = (
    "Answer only from the client project documents. Generate a high-level "
    "WBS for the demolition and site clearance scope in this project's BOQ."
)

_BOQ_ITEMS = [
    {"item_key": "D110", "description": "General site clearance", "source": "S2.md"},
    {"item_key": "D290.1", "description": "Removal of trees in existing sidewalks"},
]


def _f1_wbs() -> dict:
    acts = activities_from_boq_scope_outline(_BOQ_ITEMS)
    # Apply an explicit duration override so the workbook Dur column is not
    # the template default — H2 ground truth: durations/overrides present.
    for a in acts:
        if "D110" in (a.get("name") or ""):
            a["duration_days"] = 6
    return {
        "status": "success",
        "brief": F1_ASK,
        "actual_count": len(acts),
        "activities": acts,
        "wbs_tree": wbs_tree_from_boq_items(_BOQ_ITEMS),
        "start_date": "2026-03-01",
        "summary": {"total_duration_days": 7, "critical_count": 2},
        "duration_overrides_applied": [
            {"match": "site clearance", "days": 6, "activities_updated": 1},
        ],
        "scaffold": {
            "source": "boq",
            "derived_from_boq": True,
            "declaration": (
                "High-level WBS derived from this project's BOQ "
                "demolition and site clearance items."
            ),
        },
    }


def _scaffold_wbs() -> dict:
    return {
        "status": "success",
        "brief": F1_ASK,
        "actual_count": 204,
        "activities": [
            {"id": f"A{i}", "name": f"Hall {i}", "duration_days": 5,
             "wbs_phase": "Superstructure", "predecessors": []}
            for i in range(204)
        ],
        "wbs_tree": {"1": "Site Preparation", "2": "Superstructure"},
        "scaffold": {
            "source": "template",
            "derived_from_boq": False,
            "declaration": (
                "Template scaffold, project_type inferred: building "
                "— not derived from this project's BOQ"
            ),
        },
    }


# --------------------------------------------------------------------------
# Intent
# --------------------------------------------------------------------------


def test_verbatim_h2_ask_is_a_wbs_export():
    assert message_wants_wbs_export(H2_EXPORT_ASK)
    assert message_wants_wbs_export("Export F1 WBS as xlsx")
    assert message_wants_wbs_export("download the WBS as excel")
    assert message_wants_wbs_export("export this schedule as xlsx")


@pytest.mark.parametrize("msg", [
    F1_ASK,
    "produce the schedule",
    "Generate a high-level WBS for a 10-storey tower",
    "export the BOQ to Excel",
    "Export A1-A9 answers as a docx report",
    "how long is procurement on the schedule?",
])
def test_build_and_lookups_are_not_wbs_exports(msg):
    assert not message_wants_wbs_export(msg), msg


def test_f1_ask_still_elects_boq_scope_wbs():
    assert message_wants_boq_scope_wbs(F1_ASK)
    assert not message_wants_boq_scope_wbs(H2_EXPORT_ASK)


def test_wbs_answer_markers():
    assert message_looks_like_wbs_answer(
        "Schedule built: 2 activities over 2 working days (2 on the critical path)."
    )
    assert message_looks_like_wbs_answer("## High-level WBS\n\n### 1 Demolition")
    assert not message_looks_like_wbs_answer("The Accepted Contract Amount is SAR 1.")


# --------------------------------------------------------------------------
# Staging
# --------------------------------------------------------------------------


def test_stage_and_load_round_trip(tmp_path, monkeypatch):
    monkeypatch.setenv("DATA_DIR", str(tmp_path))
    cid = "ws-proj-h2"
    assert load_conversation_wbs(cid) is None
    snap = stage_conversation_wbs(cid, _f1_wbs())
    assert snap and len(snap["activities"]) == 2
    loaded = load_conversation_wbs(cid)
    assert loaded is not None
    names = " ".join(a["name"] for a in loaded["activities"])
    assert "D110" in names and "D290" in names
    assert loaded["scaffold"]["derived_from_boq"] is True
    assert loaded["duration_overrides_applied"][0]["days"] == 6
    assert clear_conversation_wbs(cid) is True
    assert load_conversation_wbs(cid) is None


def test_does_not_stage_a_boq_ask_that_fell_to_scaffold(tmp_path, monkeypatch):
    monkeypatch.setenv("DATA_DIR", str(tmp_path))
    cid = "ws-proj-scaffold"
    assert stage_conversation_wbs(cid, _scaffold_wbs()) is None
    assert load_conversation_wbs(cid) is None


def test_refuse_scaffold_when_boq_wbs_was_requested():
    assert refuse_scaffold_for_boq_wbs_ask(F1_ASK, F1_ASK, _scaffold_wbs())
    assert refuse_scaffold_for_boq_wbs_ask(F1_ASK, F1_ASK, _f1_wbs()) is None
    assert refuse_scaffold_for_boq_wbs_ask(
        "produce the schedule", "produce the schedule", _scaffold_wbs(),
    ) is None


def test_snapshot_requires_activities():
    assert snapshot_from_wbs({"wbs_tree": {"1": "X"}}) is None
    assert snapshot_from_wbs({"activities": []}) is None


# --------------------------------------------------------------------------
# Fulfill (chat short-circuit)
# --------------------------------------------------------------------------


def test_fulfill_without_staged_wbs_is_honest(tmp_path, monkeypatch):
    monkeypatch.setenv("DATA_DIR", str(tmp_path))
    from app.core import agent_memory
    cid = "ws-h2-empty"
    agent_memory.get_or_create_conversation(cid, "project-assistant", "p1")
    answer, exports = fulfill_wbs_export(H2_EXPORT_ASK, "p1", cid, "project-assistant")
    assert "will not substitute a generic template scaffold" in answer
    assert exports == []
    roles = [m["role"] for m in agent_memory.get_messages(cid)]
    assert roles[-2:] == ["user", "assistant"]


def test_fulfill_with_staged_f1_wbs_offers_conversation_export(tmp_path, monkeypatch):
    monkeypatch.setenv("DATA_DIR", str(tmp_path))
    from app.core import agent_memory
    cid = "ws-h2-f1"
    agent_memory.get_or_create_conversation(cid, "project-assistant", "p1")
    stage_conversation_wbs(cid, _f1_wbs())
    answer, exports = fulfill_wbs_export(H2_EXPORT_ASK, "p1", cid, "project-assistant")
    assert "2 activities" in answer
    assert "template scaffold" not in answer.lower() or "not a generic" in answer.lower()
    assert "BOQ-derived" in answer
    assert "6 days" in answer
    assert len(exports) == 1
    assert exports[0]["endpoint"].endswith(f"/conversations/{cid}/export/schedule")
    desc = conversation_schedule_export_descriptor("p1", cid, 2)
    assert desc["payload"]["conversation_id"] == cid


# --------------------------------------------------------------------------
# HTTP: conversation schedule + message xlsx
# --------------------------------------------------------------------------


@pytest.fixture
def client():
    with TestClient(app) as c:
        yield c


def _pid(client, name="H2 WBS Export"):
    r = client.post("/v1/projects", json={"name": name, "client": "ACME"}, headers=H)
    assert r.status_code == 201, r.text
    return r.json()["id"]


def _l2_activity_names(content: bytes) -> list[str]:
    wb = openpyxl.load_workbook(io.BytesIO(content))
    assert "L2 Schedule" in wb.sheetnames
    ws = wb["L2 Schedule"]
    names = []
    for row in ws.iter_rows(min_row=4, max_col=4, values_only=True):
        name, dur = row[2], row[3]
        if not name or name == "TOTAL" or str(name).startswith("Project Duration"):
            continue
        names.append((str(name), dur))
    return names


@requires_construction_kit
def test_run_workflow_stages_f1_wbs_and_binds_export_offer(tmp_path, monkeypatch):
    monkeypatch.setenv("DATA_DIR", str(tmp_path))
    from app.core.predefined_reasoning import run_workflow
    from app.schemas.project_session import ProjectSession

    cid = "ws-run-wf-f1"
    out = asyncio.run(run_workflow(
        "generate_wbs",
        {
            "message": F1_ASK,
            "project_id": "proj1",
            "conversation_id": cid,
            "boq_items": _BOQ_ITEMS,
            "params": {"target_count": 40, "boq_items": _BOQ_ITEMS},
            "deliverable": True,
        },
        ProjectSession.new("s-f1"),
    ))
    assert out["handled"] and out["status"] == "success"
    staged = load_conversation_wbs(cid)
    assert staged is not None
    assert len(staged["activities"]) == 2
    assert staged["scaffold"]["derived_from_boq"] is True
    assert out["exports"]
    assert out["exports"][0]["endpoint"].endswith(
        f"/conversations/{cid}/export/schedule"
    )
    assert "204" not in (out.get("answer") or "")


@requires_construction_kit
def test_conversation_schedule_export_is_the_staged_f1_wbs(client, tmp_path, monkeypatch):
    monkeypatch.setenv("DATA_DIR", str(tmp_path))
    pid = _pid(client)
    cid = f"ws-{pid}"
    stage_conversation_wbs(cid, _f1_wbs())

    r = client.post(
        f"/v1/projects/{pid}/conversations/{cid}/export/schedule",
        headers=H,
    )
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == _XLSX_MEDIA
    assert r.headers.get("X-Activities") == "2"
    assert r.headers.get("X-WBS-Source") == "conversation"
    assert r.content[:2] == b"PK"
    rows = _l2_activity_names(r.content)
    assert len(rows) == 2
    blob = " ".join(n for n, _ in rows)
    assert "D110" in blob and "D290" in blob
    assert "Hall" not in blob and "Superstructure" not in blob
    durs = [d for _, d in rows]
    assert 6 in durs
    assert all(d is not None and d != "" for d in durs)


@requires_construction_kit
def test_conversation_schedule_export_404_when_nothing_staged(client, tmp_path, monkeypatch):
    monkeypatch.setenv("DATA_DIR", str(tmp_path))
    pid = _pid(client)
    cid = f"ws-{pid}"
    r = client.post(
        f"/v1/projects/{pid}/conversations/{cid}/export/schedule",
        headers=H,
    )
    assert r.status_code == 404, r.text
    assert "template" in r.json()["detail"].lower()


@requires_construction_kit
def test_schedule_from_brief_uses_staged_conversation_wbs(client, tmp_path, monkeypatch):
    monkeypatch.setenv("DATA_DIR", str(tmp_path))
    pid = _pid(client, "Master Corpus")
    cid = f"ws-{pid}"
    stage_conversation_wbs(cid, _f1_wbs())

    r = client.post(
        f"/v1/projects/{pid}/export/schedule-from-brief",
        json={"brief": H2_EXPORT_ASK, "target_count": 200, "conversation_id": cid},
        headers=H,
    )
    assert r.status_code == 200, r.text
    assert r.headers.get("X-Activities") == "2"
    rows = _l2_activity_names(r.content)
    assert len(rows) == 2
    assert "D110" in " ".join(n for n, _ in rows)


@requires_construction_kit
def test_schedule_from_brief_refuses_scaffold_for_f1_without_rows(client):
    """BOQ-derived WBS asked, no rows retrieved → 422, not a 204 success."""
    pid = _pid(client)
    r = client.post(
        f"/v1/projects/{pid}/export/schedule-from-brief",
        json={"brief": F1_ASK, "target_count": 200, "project_type": "building"},
        headers=H,
    )
    assert r.status_code == 422, r.text
    detail = r.json()["detail"]
    assert "scaffold" in detail.lower()
    assert "BOQ" in detail


@requires_construction_kit
def test_per_answer_xlsx_finds_predefined_assistant_message(client, tmp_path, monkeypatch):
    """The H2 404: predefined WBS turns never wrote agent_memory.

    After a predefined generate_wbs turn the conversation must have an
    assistant message so per-answer Excel download is 200, not
    'No assistant messages in this conversation.'
    """
    monkeypatch.setenv("DATA_DIR", str(tmp_path))
    monkeypatch.setenv("ORCHESTRATOR_PREDEFINED", "1")
    pid = _pid(client, "Pred Persist")
    cid = f"ws-{pid}"

    import app.routers.chat as chat_mod

    async def run():
        evs = []
        gen = chat_mod._stream_from_predefined(
            action="generate_wbs",
            user_message="produce the schedule",
            project_id=pid,
            user_id="u1",
            session_id=cid,
            document_ids=[],
            conversation_id=cid,
        )
        async for chunk in gen:
            for line in chunk.splitlines():
                if line.startswith("data: "):
                    evs.append(json.loads(line[6:]))
        return evs

    evs = asyncio.run(run())
    assert any(e.get("type") == "end" for e in evs)

    from app.core import agent_memory
    msgs = agent_memory.get_messages(cid)
    assistant = [m for m in msgs if m.get("role") == "assistant"]
    assert assistant, "predefined turn must persist the assistant answer"
    assert "Schedule built" in (assistant[-1].get("content") or "")

    r = client.post(
        f"/v1/projects/{pid}/conversations/{cid}/export?format=xlsx&message_index=-1",
        headers=H,
    )
    assert r.status_code == 200, r.text
    assert "spreadsheetml" in r.headers["content-type"]
    assert r.content[:2] == b"PK"


@requires_construction_kit
def test_per_answer_xlsx_of_wbs_message_is_the_staged_workbook(client, tmp_path, monkeypatch):
    monkeypatch.setenv("DATA_DIR", str(tmp_path))
    from app.core import agent_memory
    pid = _pid(client)
    cid = f"ws-{pid}"
    agent_memory.get_or_create_conversation(cid, "project-assistant", pid)
    agent_memory.append_message(cid, "user", F1_ASK)
    agent_memory.append_message(
        cid, "assistant",
        "Schedule built: 2 activities over 7 working days (2 on the critical path).",
    )
    stage_conversation_wbs(cid, _f1_wbs())

    r = client.post(
        f"/v1/projects/{pid}/conversations/{cid}/export?format=xlsx&message_index=0",
        headers=H,
    )
    assert r.status_code == 200, r.text
    rows = _l2_activity_names(r.content)
    assert len(rows) == 2
    assert "D110" in " ".join(n for n, _ in rows)
