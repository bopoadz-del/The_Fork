"""In-place re-extract of stale TEXT_SPARSE .docx — no add_document, no crawl."""
from __future__ import annotations

import ast
import hashlib
import importlib
import io
from pathlib import Path

import docx
from sqlalchemy import func, select

from app.core.ingest_status import EXTRACTOR_VERSION, INDEXED, TEXT_SPARSE


DRIVE_STALE_ID = "driveStale01"
FORBIDDEN_CALLS = frozenset({
    "add_document",
    "supersede_document",
    "archive_document",
})
SCRIPT_PATH = Path(__file__).resolve().parent.parent / "scripts" / "reextract_stale_docx.py"


def _reload(monkeypatch, tmp_path):
    monkeypatch.setenv("DATA_DIR", str(tmp_path))
    monkeypatch.setenv("RAG_EMBEDDING_MODEL", "fake")
    monkeypatch.delenv("DATABASE_URL", raising=False)
    import app.core.db as db_mod
    import app.core.users as users_mod
    from app.core import projects
    from app.core.rag.embeddings import reset_embedder_cache
    from app.core.rag.vector_store import reset_store_cache

    importlib.reload(db_mod)
    importlib.reload(users_mod)
    users_mod._initialized = False
    projects._initialized = False
    reset_store_cache()
    reset_embedder_cache()
    import app.core.doc_index as doc_index_mod
    importlib.reload(doc_index_mod)
    return importlib.reload(projects), users_mod


def _rich_docx_bytes() -> bytes:
    """Enough extracted words to clear the single-window TEXT_SPARSE gate."""
    document = docx.Document()
    document.add_paragraph(" ".join(f"word{i}" for i in range(800)))
    document.add_paragraph(" ".join(f"extra{i}" for i in range(200)))
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def _thin_docx_bytes() -> bytes:
    document = docx.Document()
    document.add_paragraph("one window only")
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def _seed_project(projects, users):
    projects.init_db()
    users.ensure_user_exists("u1")
    projects.create_project(name="P", client="C", user_id="u1")
    return projects.list_projects("u1")[0]


def _add_docx(
    projects,
    project_id: str,
    tmp_path: Path,
    *,
    name: str,
    ingest_status: str,
    extractor_version: str | None,
    chunk_count: int,
    drive_file_id: str = "",
    r2_object_key: str | None = "projects/p/stale.docx",
    retrieval_visible: bool = True,
):
    dest = tmp_path / name
    raw = _thin_docx_bytes()
    dest.write_bytes(raw)
    metadata: dict = {}
    if drive_file_id:
        metadata["drive_file_id"] = drive_file_id
    if r2_object_key:
        metadata["r2_object_key"] = r2_object_key
        metadata["r2_bucket"] = "corpus"
    doc = projects.add_document(
        project_id=project_id,
        original_name=name,
        stored_as=name,
        file_path=str(dest),
        size=len(raw),
        content_sha256=hashlib.sha256(name.encode()).hexdigest(),
        metadata=metadata,
    )
    projects.stamp_document_index(
        doc["id"],
        chunk_count=chunk_count,
        ingest_status=ingest_status,
        ingest_status_reason="single_window:terminal",
        extractor_version=extractor_version,
    )
    if not retrieval_visible:
        from app.core.db import SessionLocal
        from app.core.models import Document

        with SessionLocal() as session:
            row = session.get(Document, doc["id"])
            row.retrieval_visible = False
            session.commit()
    return projects.get_document(doc["id"])


def _document_count() -> int:
    from app.core.db import SessionLocal
    from app.core.models import Document

    with SessionLocal() as session:
        return int(session.scalar(select(func.count()).select_from(Document)) or 0)


def test_script_forbids_row_creating_and_reingest_paths():
    source = SCRIPT_PATH.read_text(encoding="utf-8")
    tree = ast.parse(source)
    called: set[str] = set()
    for node in ast.walk(tree):
        if not isinstance(node, ast.Call):
            continue
        func = node.func
        if isinstance(func, ast.Name):
            called.add(func.id)
        elif isinstance(func, ast.Attribute):
            called.add(func.attr)
    assert not (called & FORBIDDEN_CALLS)
    assert "--reingest" not in source


def test_selection_picks_exactly_the_stale_docx(monkeypatch, tmp_path):
    projects, users = _reload(monkeypatch, tmp_path)
    proj = _seed_project(projects, users)
    stale = _add_docx(
        projects, proj["id"], tmp_path,
        name="stale.docx",
        ingest_status=TEXT_SPARSE,
        extractor_version="pre-sdt",
        chunk_count=1,
        drive_file_id=DRIVE_STALE_ID,
    )
    _add_docx(
        projects, proj["id"], tmp_path,
        name="current.docx",
        ingest_status=TEXT_SPARSE,
        extractor_version=EXTRACTOR_VERSION,
        chunk_count=1,
        drive_file_id="driveCurre01",
        r2_object_key="projects/p/current.docx",
    )
    _add_docx(
        projects, proj["id"], tmp_path,
        name="hidden.docx",
        ingest_status=TEXT_SPARSE,
        extractor_version="pre-sdt",
        chunk_count=1,
        drive_file_id="driveHidde01",
        r2_object_key="projects/p/hidden.docx",
        retrieval_visible=False,
    )
    pdf = projects.add_document(
        project_id=proj["id"],
        original_name="sheet.pdf",
        stored_as="sheet.pdf",
        file_path=str(tmp_path / "sheet.pdf"),
        size=10,
        content_sha256="aa" * 32,
        metadata={"drive_file_id": "drivePdfxx01"},
    )
    projects.stamp_document_index(
        pdf["id"],
        chunk_count=1,
        ingest_status=TEXT_SPARSE,
        ingest_status_reason="single_window:terminal",
        extractor_version="pre-sdt",
    )

    from scripts.reextract_stale_docx import select_stale_docx_rows

    selected = select_stale_docx_rows()
    assert [row["id"] for row in selected] == [stale["id"]]


def test_reextract_same_id_indexed_row_count_unchanged(monkeypatch, tmp_path, capsys):
    projects, users = _reload(monkeypatch, tmp_path)
    proj = _seed_project(projects, users)
    stale = _add_docx(
        projects, proj["id"], tmp_path,
        name="stale.docx",
        ingest_status=TEXT_SPARSE,
        extractor_version="pre-sdt",
        chunk_count=1,
        drive_file_id=DRIVE_STALE_ID,
    )
    current = _add_docx(
        projects, proj["id"], tmp_path,
        name="current.docx",
        ingest_status=INDEXED,
        extractor_version=EXTRACTOR_VERSION,
        chunk_count=5,
        drive_file_id="driveCurre01",
        r2_object_key="projects/p/current.docx",
    )
    before_count = _document_count()
    rich = _rich_docx_bytes()

    monkeypatch.setattr(
        "app.core.r2_storage.fetch_object_bytes",
        lambda key, bucket=None: rich if key == "projects/p/stale.docx" else None,
    )
    drive_calls: list[str] = []
    monkeypatch.setattr(
        "app.core.gdrive_service.download_file_bytes",
        lambda fid: drive_calls.append(fid) or (None, "unused"),
    )

    from scripts.reextract_stale_docx import main

    assert main([]) == 0
    out = capsys.readouterr().out
    after = projects.get_document(stale["id"])
    assert after["id"] == stale["id"]
    assert after["chunk_count"] > 1
    assert after["ingest_status"] == INDEXED
    assert after["extractor_version"] == EXTRACTOR_VERSION
    assert _document_count() == before_count
    assert projects.get_document(current["id"])["extractor_version"] == EXTRACTOR_VERSION
    assert f"VERIFICATION doc_id={stale['id']}" in out
    assert "before=TEXT_SPARSE/pre-sdt/1" in out
    assert f"after={INDEXED}/{EXTRACTOR_VERSION}/" in out
    assert drive_calls == []


def test_r2_none_falls_back_to_one_drive_file_never_folder(
    monkeypatch, tmp_path, capsys,
):
    projects, users = _reload(monkeypatch, tmp_path)
    proj = _seed_project(projects, users)
    stale = _add_docx(
        projects, proj["id"], tmp_path,
        name="stale.docx",
        ingest_status=TEXT_SPARSE,
        extractor_version="pre-sdt",
        chunk_count=1,
        drive_file_id=DRIVE_STALE_ID,
    )
    rich = _rich_docx_bytes()
    drive_calls: list[str] = []
    folder_calls: list[str] = []

    monkeypatch.setattr(
        "app.core.r2_storage.fetch_object_bytes",
        lambda key, bucket=None: None,
    )
    monkeypatch.setattr(
        "app.core.gdrive_service.download_file_bytes",
        lambda fid: drive_calls.append(fid) or (rich, None),
    )
    monkeypatch.setattr(
        "app.core.gdrive_service.walk_folder",
        lambda *a, **k: folder_calls.append("walk_folder") or ([], []),
    )
    monkeypatch.setattr(
        "app.core.gdrive_service.list_folder_files",
        lambda *a, **k: folder_calls.append("list_folder_files") or ([], None),
    )

    from scripts.reextract_stale_docx import main

    assert main([]) == 0
    assert drive_calls == [DRIVE_STALE_ID]
    assert folder_calls == []
    after = projects.get_document(stale["id"])
    assert after["id"] == stale["id"]
    assert after["ingest_status"] == INDEXED
    assert after["extractor_version"] == EXTRACTOR_VERSION
    assert after["chunk_count"] > 1
    assert "SOURCE_UNAVAILABLE" not in capsys.readouterr().out


def test_both_sources_none_source_unavailable_row_untouched_exit_1(
    monkeypatch, tmp_path, capsys,
):
    projects, users = _reload(monkeypatch, tmp_path)
    proj = _seed_project(projects, users)
    stale = _add_docx(
        projects, proj["id"], tmp_path,
        name="stale.docx",
        ingest_status=TEXT_SPARSE,
        extractor_version="pre-sdt",
        chunk_count=1,
        drive_file_id=DRIVE_STALE_ID,
    )
    current = _add_docx(
        projects, proj["id"], tmp_path,
        name="current.docx",
        ingest_status=INDEXED,
        extractor_version=EXTRACTOR_VERSION,
        chunk_count=5,
        drive_file_id="driveCurre01",
        r2_object_key="projects/p/current.docx",
    )
    original = Path(stale["file_path"]).read_bytes()
    before_count = _document_count()

    monkeypatch.setattr(
        "app.core.r2_storage.fetch_object_bytes",
        lambda key, bucket=None: None,
    )
    monkeypatch.setattr(
        "app.core.gdrive_service.download_file_bytes",
        lambda fid: (None, "missing"),
    )

    from scripts.reextract_stale_docx import main

    assert main([]) == 1
    out = capsys.readouterr().out
    after = projects.get_document(stale["id"])
    assert after["id"] == stale["id"]
    assert after["ingest_status"] == TEXT_SPARSE
    assert after["extractor_version"] == "pre-sdt"
    assert after["chunk_count"] == 1
    assert Path(stale["file_path"]).read_bytes() == original
    assert _document_count() == before_count
    assert projects.get_document(current["id"])["chunk_count"] == 5
    assert "source_unavailable=1" in out
    assert f"SOURCE_UNAVAILABLE doc_id={stale['id']} drive_file_id={DRIVE_STALE_ID}" in out


def test_dry_run_prints_resolved_source_and_writes_nothing(monkeypatch, tmp_path, capsys):
    projects, users = _reload(monkeypatch, tmp_path)
    proj = _seed_project(projects, users)
    stale = _add_docx(
        projects, proj["id"], tmp_path,
        name="stale.docx",
        ingest_status=TEXT_SPARSE,
        extractor_version="pre-sdt",
        chunk_count=1,
        drive_file_id=DRIVE_STALE_ID,
    )
    original = Path(stale["file_path"]).read_bytes()
    writes: list[str] = []
    indexes: list[str] = []

    monkeypatch.setattr(
        "app.core.r2_storage.fetch_object_bytes",
        lambda key, bucket=None: b"docx-bytes",
    )
    monkeypatch.setattr(
        "app.core.file_crypto.write_document",
        lambda path, data: writes.append(path),
    )
    monkeypatch.setattr(
        "app.core.doc_index.index_document",
        lambda *a, **k: indexes.append("indexed") or {"status": "ok"},
    )

    from scripts.reextract_stale_docx import main

    assert main(["--dry-run"]) == 0
    out = capsys.readouterr().out
    after = projects.get_document(stale["id"])
    assert after["ingest_status"] == TEXT_SPARSE
    assert after["extractor_version"] == "pre-sdt"
    assert Path(stale["file_path"]).read_bytes() == original
    assert writes == []
    assert indexes == []
    assert f"doc_id={stale['id']}" in out
    assert f"drive_file_id={DRIVE_STALE_ID}" in out
    assert "source=r2" in out
    assert "source_unavailable=0" in out
    assert "would_write=" not in out


def _seed_leading_none_then_fetchable(projects, users, tmp_path):
    """Two pointer-less open rows, then three R2-backed stale rows."""
    proj = _seed_project(projects, users)
    none_rows = []
    for name in ("none_a.docx", "none_b.docx"):
        none_rows.append(
            _add_docx(
                projects, proj["id"], tmp_path,
                name=name,
                ingest_status=TEXT_SPARSE,
                extractor_version="pre-sdt",
                chunk_count=1,
                drive_file_id="",
                r2_object_key=None,
            )
        )
    fetchable = []
    for name, key in (
        ("fetch_a.docx", "projects/p/fetch_a.docx"),
        ("fetch_b.docx", "projects/p/fetch_b.docx"),
        ("fetch_c.docx", "projects/p/fetch_c.docx"),
    ):
        fetchable.append(
            _add_docx(
                projects, proj["id"], tmp_path,
                name=name,
                ingest_status=TEXT_SPARSE,
                extractor_version="pre-sdt",
                chunk_count=1,
                drive_file_id=f"drive{name[:6]}01",
                r2_object_key=key,
            )
        )
    return none_rows, fetchable


def test_limit_skips_source_none_and_retries_fetchable(
    monkeypatch, tmp_path, capsys,
):
    projects, users = _reload(monkeypatch, tmp_path)
    none_rows, fetchable = _seed_leading_none_then_fetchable(
        projects, users, tmp_path,
    )
    rich = _rich_docx_bytes()

    monkeypatch.setattr(
        "app.core.r2_storage.fetch_object_bytes",
        lambda key, bucket=None: (
            rich if key in {
                "projects/p/fetch_a.docx",
                "projects/p/fetch_b.docx",
                "projects/p/fetch_c.docx",
            } else None
        ),
    )
    monkeypatch.setattr(
        "app.core.gdrive_service.download_file_bytes",
        lambda fid: (None, "unused"),
    )

    from scripts.reextract_stale_docx import main, select_stale_docx_rows

    fetch_ids = {row["id"] for row in fetchable}
    ordered_fetchable = [
        row["id"] for row in select_stale_docx_rows() if row["id"] in fetch_ids
    ]
    write_ids = set(ordered_fetchable[:2])
    leftover_id = ordered_fetchable[2]

    assert main(["--limit", "2"]) == 1
    out = capsys.readouterr().out
    for doc_id in write_ids:
        after = projects.get_document(doc_id)
        assert after["ingest_status"] == INDEXED
        assert after["extractor_version"] == EXTRACTOR_VERSION
    leftover = projects.get_document(leftover_id)
    assert leftover["ingest_status"] == TEXT_SPARSE
    assert leftover["extractor_version"] == "pre-sdt"
    after_none = projects.get_document(none_rows[0]["id"])
    assert after_none["ingest_status"] == TEXT_SPARSE
    assert after_none["extractor_version"] == "pre-sdt"
    assert "open=5" in out
    assert "retried=2" in out
    assert "source_unavailable=2" in out
    assert f"SOURCE_UNAVAILABLE doc_id={none_rows[0]['id']}" in out
    assert f"SOURCE_UNAVAILABLE doc_id={none_rows[1]['id']}" in out
    for doc_id in write_ids:
        assert f"VERIFICATION doc_id={doc_id}" in out
    assert f"after={INDEXED}/{EXTRACTOR_VERSION}/" in out


def test_dry_run_limit_prints_all_open_and_counts_unavailable(
    monkeypatch, tmp_path, capsys,
):
    projects, users = _reload(monkeypatch, tmp_path)
    none_rows, fetchable = _seed_leading_none_then_fetchable(
        projects, users, tmp_path,
    )
    writes: list[str] = []
    indexes: list[str] = []
    monkeypatch.setattr(
        "app.core.r2_storage.fetch_object_bytes",
        lambda key, bucket=None: (
            b"docx-bytes" if key in {
                "projects/p/fetch_a.docx",
                "projects/p/fetch_b.docx",
                "projects/p/fetch_c.docx",
            } else None
        ),
    )
    monkeypatch.setattr(
        "app.core.file_crypto.write_document",
        lambda path, data: writes.append(path),
    )
    monkeypatch.setattr(
        "app.core.doc_index.index_document",
        lambda *a, **k: indexes.append("indexed") or {"status": "ok"},
    )

    from scripts.reextract_stale_docx import main

    assert main(["--dry-run", "--limit", "2"]) == 0
    out = capsys.readouterr().out
    assert writes == []
    assert indexes == []
    for row in none_rows + fetchable:
        after = projects.get_document(row["id"])
        assert after["ingest_status"] == TEXT_SPARSE
        assert f"DRY-RUN doc_id={row['id']}" in out
    assert f"DRY-RUN doc_id={none_rows[0]['id']}" in out
    assert "source=none would_write=0" in out
    assert "source=r2 would_write=1" in out
    assert out.count("would_write=1") == 2
    assert out.count("would_write=0") == 3
    assert "open=5" in out
    assert "source_unavailable=2" in out
    assert "dry_run=1" in out


def test_allow_unavailable_exits_0_when_only_none_remain(
    monkeypatch, tmp_path, capsys,
):
    projects, users = _reload(monkeypatch, tmp_path)
    none_rows, fetchable = _seed_leading_none_then_fetchable(
        projects, users, tmp_path,
    )
    rich = _rich_docx_bytes()
    monkeypatch.setattr(
        "app.core.r2_storage.fetch_object_bytes",
        lambda key, bucket=None: (
            rich if key in {
                "projects/p/fetch_a.docx",
                "projects/p/fetch_b.docx",
                "projects/p/fetch_c.docx",
            } else None
        ),
    )
    monkeypatch.setattr(
        "app.core.gdrive_service.download_file_bytes",
        lambda fid: (None, "unused"),
    )

    from scripts.reextract_stale_docx import main, select_stale_docx_rows

    fetch_ids = {row["id"] for row in fetchable}
    ordered_fetchable = [
        row["id"] for row in select_stale_docx_rows() if row["id"] in fetch_ids
    ]
    write_ids = set(ordered_fetchable[:2])
    leftover_id = ordered_fetchable[2]

    assert main(["--limit", "2", "--allow-unavailable"]) == 0
    out = capsys.readouterr().out
    for doc_id in write_ids:
        assert projects.get_document(doc_id)["ingest_status"] == INDEXED
    assert projects.get_document(leftover_id)["ingest_status"] == TEXT_SPARSE
    assert projects.get_document(none_rows[0]["id"])["ingest_status"] == TEXT_SPARSE
    assert "retried=2" in out
    assert "source_unavailable=2" in out
